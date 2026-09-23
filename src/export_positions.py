"""Полная выгрузка позиций прайсов с направлениями (392 917 строк).

Источники: data/prices.db (позиции, флаги цен), data/directions.db (направление
названия), data/osint.db (компания, город, регион по ИНН). Ничего не
достраивается: неразмеченное название выводится как «Не классифицировано».
"""

import collections
import datetime
import re
import sqlite3

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.direction_map import norm, strip_ui


_ILLEGAL = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")


def clean(v):
    """Excel не принимает управляющие символы, встречающиеся в сырых прайсах."""
    return _ILLEGAL.sub(" ", v) if isinstance(v, str) else v


def companies():
    db = sqlite3.connect("file:data/osint.db?mode=ro", uri=True)
    out = {}
    for t in ("t40_companies", "companies"):
        try:
            for inn, name, city, region in db.execute(
                    f"select inn, name, city, region from {t}"):
                out.setdefault(str(inn), (name, city, region))
        except sqlite3.OperationalError:
            continue
    return out


def build(path=None):
    path = path or f"output/Позиции_прайсов_с_направлениями_{datetime.date.today()}.xlsx"
    dd = sqlite3.connect("file:data/directions.db?mode=ro", uri=True)
    dmap = {k: (d, c, m) for k, d, c, m in dd.execute(
        "select name_norm, direction, confidence, method from name_directions")}
    p = sqlite3.connect("file:data/prices.db?mode=ro", uri=True)
    flags = {r[0]: r[1] for r in p.execute("select price_item_id, flag from price_flags")}
    comp = companies()

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Позиции_с_направлениями")
    head = ["ИНН", "Компания", "Город", "Регион", "Домен", "Раздел прайса",
            "Название услуги (дословно)", "Цена ₽ (очищенная)", "Цена ₽ (исходная)",
            "Флаг цены", "Направление", "Уверенность", "Метод разметки"]
    ws.append(head)
    stat = collections.Counter()
    n = 0
    for pid, inn, domain, section, name, praw, price in p.execute(
            "select id, inn, domain, section, name_raw, price_raw, price_value "
            "from price_items order by inn"):
        d, conf, method = dmap.get(norm(strip_ui(name or "")),
                                   ("Не классифицировано", None, None))
        flag = flags.get(pid)
        c = comp.get(str(inn), (None, None, None))
        ws.append([inn, clean(c[0]), c[1], c[2], domain, clean(section), clean(name),
                   None if flag else price, clean(praw), flag, d, conf, method])
        stat[d] += 1
        n += 1
    sm = wb.create_sheet("Сводка")
    sm.append(["Направление", "Позиций", "Доля"])
    for d, k in stat.most_common():
        sm.append([d, k, round(k / n, 4)])
    wb.save(path)
    return path, n, stat


if __name__ == "__main__":
    path, n, stat = build()
    print(path, n, "строк")
    for d, k in stat.most_common(8):
        print(f"  {k:7d}  {d}")
